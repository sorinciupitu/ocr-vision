"""
OCR Service - Uses LightOnOCR-3-4B via vLLM OpenAI-compatible API.
Handles: PDF rendering, image preprocessing, DOCX text extraction,
and sending images to the model for OCR transcription.
"""
import base64
import io
import os
import logging
from pathlib import Path
from typing import List, Tuple, Optional

import httpx
from PIL import Image
from docx import Document as DocxDocument

from app.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()


async def render_pdf_pages(pdf_path: str) -> List[Image.Image]:
    """Render each page of a PDF to a PIL Image at 400 DPI, thumbnail to 2048px."""
    import pypdfium2 as pdfium

    images = []
    pdf = pdfium.PdfDocument(pdf_path)
    for i in range(len(pdf)):
        page = pdf[i]
        bitmap = page.render(scale=400 / 72)
        pil_image = bitmap.to_pil()
        pil_image.thumbnail((2048, 2048))
        images.append(pil_image)
    return images


def image_to_base64(img: Image.Image) -> str:
    """Convert PIL Image to base64 PNG string."""
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode("utf-8")


def load_image(path: str) -> Image.Image:
    """Load an image file and resize to max 2048px."""
    img = Image.open(path)
    if img.mode != "RGB":
        img = img.convert("RGB")
    img.thumbnail((2048, 2048))
    return img


def extract_docx_text(path: str) -> str:
    """Extract plain text from a DOCX file."""
    doc = DocxDocument(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


async def ocr_image(img: Image.Image, prompt: str = "") -> str:
    """Send a single image to LightOnOCR via vLLM API and return the transcribed text."""
    b64 = image_to_base64(img)

    content = [
        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}}
    ]
    if prompt:
        content.append({"type": "text", "text": prompt})

    payload = {
        "model": settings.vllm_model,
        "messages": [{"role": "user", "content": content}],
        "max_tokens": 4096,
        "temperature": 0.2,
        "top_p": 0.9,
    }

    async with httpx.AsyncClient(timeout=120.0) as client:
        try:
            response = await client.post(settings.vllm_url, json=payload)
            response.raise_for_status()
            result = response.json()
            return result["choices"][0]["message"]["content"]
        except httpx.HTTPError as e:
            logger.error(f"OCR API error: {e}")
            raise RuntimeError(f"OCR API request failed: {e}")


async def process_file(file_path: str) -> Tuple[str, int, List[str]]:
    """
    Process a file through OCR pipeline.
    Returns: (full_text, page_count, list_of_page_texts)
    """
    ext = Path(file_path).suffix.lower()

    if ext == ".docx":
        text = extract_docx_text(file_path)
        return text, 1, [text]

    if ext == ".pdf":
        images = await render_pdf_pages(file_path)
        page_texts = []
        for img in images:
            page_text = await ocr_image(img)
            page_texts.append(page_text)
        full_text = "\n\n---PAGE BREAK---\n\n".join(page_texts)
        return full_text, len(images), page_texts

    if ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"):
        img = load_image(file_path)
        text = await ocr_image(img)
        return text, 1, [text]

    raise ValueError(f"Unsupported file type: {ext}")


def create_thumbnail(file_path: str, output_path: str, size: Tuple[int, int] = (300, 400)):
    """Create a thumbnail for a document."""
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(file_path)
        if len(pdf) > 0:
            page = pdf[0]
            bitmap = page.render(scale=150 / 72)
            img = bitmap.to_pil()
            img.thumbnail(size)
            img.save(output_path, format="PNG")
            return output_path
    elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"):
        img = Image.open(file_path)
        img.thumbnail(size)
        img.save(output_path, format="PNG")
        return output_path

    return None
